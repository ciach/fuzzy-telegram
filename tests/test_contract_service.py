from __future__ import annotations

import io

import pytest
from docx import Document

from backend.services.contract_service import (
    ContractService,
    ManualClauseLockedError,
    MutationValidationError,
)
from backend.services.logging_service import AuditLoggingService
from backend.services.rendering_service import RenderingService
from backend.storage.in_memory_store import InMemoryStore


def build_service() -> ContractService:
    return ContractService(store=InMemoryStore(), audit_logging_service=AuditLoggingService())


def test_financial_remaining_amount_is_derived_and_logged() -> None:
    service = build_service()
    contract = service.create_contract(language="es", location="Barcelona")

    service.apply_update(
        contract_id=contract.id,
        path="financial.total_price",
        value=100000,
        action_type="update_field",
        triggered_by="user",
    )
    service.apply_update(
        contract_id=contract.id,
        path="financial.arras_amount",
        value=10000,
        action_type="update_field",
        triggered_by="user",
    )

    schema = service.get_contract_schema(contract.id)
    assert schema.financial.remaining_amount == 90000

    logs = service.get_logs(contract.id)
    assert len(logs) == 2
    assert logs[-1].field_path == "financial.arras_amount"


def test_invalid_path_is_rejected() -> None:
    service = build_service()
    contract = service.create_contract(language="es", location="Barcelona")

    with pytest.raises(MutationValidationError):
        service.apply_update(
            contract_id=contract.id,
            path="financial.missing_field",
            value=10,
            action_type="update_field",
            triggered_by="user",
        )


def test_manual_clause_lock_prevents_ai_override_without_flag() -> None:
    service = build_service()
    contract = service.create_contract(language="es", location="Barcelona")

    service.rewrite_clause(
        contract_id=contract.id,
        clause_id="clause_1",
        new_text="Manual edit",
        triggered_by="user",
    )

    with pytest.raises(ManualClauseLockedError):
        service.rewrite_clause(
            contract_id=contract.id,
            clause_id="clause_1",
            new_text="AI rewrite",
            triggered_by="ai",
            allow_manual_clause_override=False,
        )


def test_ai_interaction_is_logged_even_without_mutation() -> None:
    service = build_service()
    contract = service.create_contract(language="es", location="Barcelona")

    interaction = service.log_ai_interaction(
        contract_id=contract.id,
        prompt="What should I update?",
        raw_response='{"action":"no_action","reason":"insufficient_data"}',
        parsed_action={"action": "no_action", "reason": "insufficient_data"},
        status="insufficient_data",
    )

    rows = service.get_ai_interactions(contract.id)
    assert len(rows) == 1
    assert rows[0].id == interaction.id
    assert rows[0].status == "insufficient_data"


def test_new_contract_is_seeded_with_reference_clauses() -> None:
    service = build_service()
    contract = service.create_contract(language="es", location="Barcelona")
    schema = service.get_contract_schema(contract.id)
    assert "clause_1" in schema.clauses
    assert "clause_2" in schema.clauses
    assert "{{buyers_names}}" in schema.clauses["clause_1"]
    assert "{{sellers_names}}" in schema.clauses["clause_1"]
    assert "{{total_price_eur}}" in schema.clauses["clause_1"]
    assert "{{arras_amount_eur}}" in schema.clauses["clause_1"]
    assert "{{remaining_amount_eur}}" in schema.clauses["clause_2"]
    assert "{{deadline_date}}" in schema.clauses["clause_2"]


def test_rendering_service_exports_docx_and_pdf() -> None:
    service = build_service()
    contract = service.create_contract(language="es", location="Barcelona")
    renderer = RenderingService(contract_service=service)

    docx_bytes = renderer.export_docx(contract.id)
    pdf_bytes = renderer.export_pdf(contract.id)

    assert len(docx_bytes) > 100
    assert docx_bytes[:2] == b"PK"
    assert len(pdf_bytes) > 100
    assert pdf_bytes[:4] == b"%PDF"


def test_rendering_service_uses_placeholders_for_missing_core_fields() -> None:
    service = build_service()
    contract = service.create_contract(language="es", location="Barcelona")
    renderer = RenderingService(contract_service=service)

    docx_bytes = renderer.export_docx(contract.id)
    document = Document(io.BytesIO(docx_bytes))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "SELLER_PLACEHOLDER" in text
    assert "BUYER_PLACEHOLDER" in text
    assert "ADDRESS_PLACEHOLDER" in text
    assert "PRICE_PLACEHOLDER" in text
